"""Turn DataFrames into a polished ``.xlsx`` workbook.

Plain ``df.to_excel()`` produces a wall of unstyled cells. This module keeps
the workbook readable: branded bold headers, a frozen header row, an
autofilter, sensible column widths, and number/date formats taken from the
column dtypes.

Brand palette comes from the celist.db app icon: deep indigo ``#261A54`` for
headers with the electric-blue accent used elsewhere in the product.
"""

from __future__ import annotations

from pathlib import Path
from typing import Mapping, Optional, Union

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PathLike = Union[str, Path]

# --- brand ---------------------------------------------------------------
BRAND_INDIGO = "261A54"  # header fill, taken from the app icon

# --- tunables ------------------------------------------------------------
HEADER_HEIGHT = 20
MIN_COL_WIDTH = 9
MAX_COL_WIDTH = 60
WIDTH_SAMPLE_ROWS = 200
FORMAT_MAX_ROWS = 20_000  # skip per-cell number formats on very large sheets
INVALID_SHEET_CHARS = set('[]:*?/\\')

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor=BRAND_INDIGO)
_HEADER_ALIGN = Alignment(vertical="center", horizontal="left")


# ------------------------------------------------------------------------
# helpers
# ------------------------------------------------------------------------
def sanitize_sheet_name(name: str, default: str = "Sheet") -> str:
    """Make *name* legal for Excel: max 31 chars, no ``[]:*?/\\``, not quoted."""
    cleaned = "".join("_" if ch in INVALID_SHEET_CHARS else ch for ch in str(name))
    cleaned = cleaned.strip().strip("'")
    if not cleaned:
        cleaned = default
    return cleaned[:31]


def unique_sheet_name(name: str, taken: set[str], default: str = "Sheet") -> str:
    """Sanitize *name* and make it unique within *taken* (case-insensitive)."""
    base = sanitize_sheet_name(name, default=default)
    candidate = base
    counter = 2
    while candidate.lower() in taken:
        suffix = f" ({counter})"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    taken.add(candidate.lower())
    return candidate


def _is_missing(value) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _display(value) -> str:
    return "" if _is_missing(value) else str(value)


def _hexify(value):
    """Render binary blobs as ``0x…`` hex so Excel can hold them."""
    if isinstance(value, (bytes, bytearray, memoryview)):
        text = "0x" + bytes(value).hex()
        return text if len(text) <= 66 else text[:65] + "…"
    return value


def prepare_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Copy *df* with Excel-unfriendly values converted (blobs → hex)."""
    out = df.copy()
    for column in out.columns:
        series = out[column]
        if series.dtype == object:
            if series.map(lambda v: isinstance(v, (bytes, bytearray, memoryview))).any():
                out[column] = series.map(_hexify)
    return out


def _number_format(series: pd.Series) -> Optional[str]:
    if pd.api.types.is_bool_dtype(series):
        return None
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    if pd.api.types.is_float_dtype(series):
        return "#,##0.00"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "yyyy-mm-dd hh:mm"
    return None


def _column_width(header, series: pd.Series) -> float:
    longest = len(str(header))
    for value in series.head(WIDTH_SAMPLE_ROWS):
        longest = max(longest, len(_display(value)))
    return float(min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, longest + 2)))


def _style_sheet(worksheet, df: pd.DataFrame, *, freeze_header: bool, autofilter: bool) -> None:
    """Apply headers, freeze panes, autofilter, widths, and number formats."""
    ncols = len(df.columns)
    nrows = len(df)
    if ncols == 0:
        return

    for index in range(1, ncols + 1):
        cell = worksheet.cell(row=1, column=index)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    worksheet.row_dimensions[1].height = HEADER_HEIGHT

    if freeze_header:
        worksheet.freeze_panes = "A2"
    if autofilter:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"

    apply_formats = nrows <= FORMAT_MAX_ROWS
    for index, column in enumerate(df.columns, start=1):
        series = df[column]
        worksheet.column_dimensions[get_column_letter(index)].width = _column_width(column, series)
        if not apply_formats:
            continue
        number_format = _number_format(series)
        if number_format is None:
            continue
        for row in range(2, nrows + 2):
            worksheet.cell(row=row, column=index).number_format = number_format


def _index_frame(sheets: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Sheet": list(sheets.keys()),
            "Rows": [len(df) for df in sheets.values()],
            "Columns": [len(df.columns) for df in sheets.values()],
        }
    )


# ------------------------------------------------------------------------
# public API
# ------------------------------------------------------------------------
def write_workbook(
    sheets: Mapping[str, pd.DataFrame],
    path: PathLike,
    *,
    index_sheet: bool = False,
    freeze_header: bool = True,
    autofilter: bool = True,
) -> Path:
    """Write *sheets* to *path* as one styled ``.xlsx`` workbook.

    ``index_sheet`` adds a leading "Index" sheet summarising the others — used
    by the CLI when exporting a whole database.
    """
    if not sheets:
        raise ValueError("no sheets to write")

    target = Path(path).expanduser()
    target.parent.mkdir(parents=True, exist_ok=True)

    taken: set[str] = set()
    prepared: dict[str, pd.DataFrame] = {}
    for raw_name, df in sheets.items():
        frame = prepare_dataframe(df)
        if frame.shape[1] == 0:
            frame = pd.DataFrame({"note": ["(no columns)"]})
        prepared[unique_sheet_name(raw_name, taken)] = frame

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        if index_sheet and len(prepared) > 1:
            index = _index_frame(prepared)
            index.to_excel(writer, sheet_name="Index", index=False)
            _style_sheet(
                writer.book["Index"], index, freeze_header=freeze_header, autofilter=False
            )
        for name, frame in prepared.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(
                writer.book[name], frame, freeze_header=freeze_header, autofilter=autofilter
            )

    return target
